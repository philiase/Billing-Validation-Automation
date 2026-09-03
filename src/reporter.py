from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def build_summary(
    source_rows: int,
    valid_records: pd.DataFrame,
    exceptions: pd.DataFrame,
    run_id: str,
    duration_seconds: float,
) -> pd.DataFrame:
    """Create the top-level run summary used by Billing, Finance and Operations."""
    exception_count = len(exceptions)
    unique_exception_rows = exceptions["record_id"].nunique() if exception_count else 0
    valid_count = len(valid_records)
    success_rate = valid_count / source_rows if source_rows else 0
    exception_rate = unique_exception_rows / source_rows if source_rows else 0

    return pd.DataFrame(
        [
            {"Metric": "Run ID", "Value": run_id},
            {"Metric": "Source records processed", "Value": source_rows},
            {"Metric": "Valid records", "Value": valid_count},
            {"Metric": "Rows with exceptions", "Value": unique_exception_rows},
            {"Metric": "Total exception records", "Value": exception_count},
            {"Metric": "Billing accuracy", "Value": success_rate},
            {"Metric": "Exception rate", "Value": exception_rate},
            {"Metric": "Duration seconds", "Value": round(duration_seconds, 2)},
        ]
    )


def build_exception_type_summary(exceptions: pd.DataFrame) -> pd.DataFrame:
    """Group exceptions so stakeholders can see recurring failure categories."""
    if exceptions.empty:
        return pd.DataFrame(columns=["exception_type", "severity", "exception_count"])

    return (
        exceptions.groupby(["exception_type", "severity"], dropna=False)
        .size()
        .reset_index(name="exception_count")
        .sort_values(["severity", "exception_count"], ascending=[True, False])
    )


def build_severity_summary(exceptions: pd.DataFrame) -> pd.DataFrame:
    """Show the control risk profile for the automation run."""
    if exceptions.empty:
        return pd.DataFrame(columns=["severity", "exception_count"])

    severity_order = {"Critical": 1, "High": 2, "Medium": 3, "Low": 4}
    summary = (
        exceptions.groupby("severity", dropna=False)
        .size()
        .reset_index(name="exception_count")
    )
    summary["sort_order"] = summary["severity"].map(severity_order).fillna(99)
    return summary.sort_values("sort_order").drop(columns=["sort_order"])


def create_excel_report(
    report_path: Path,
    summary: pd.DataFrame,
    exceptions: pd.DataFrame,
    valid_records: pd.DataFrame,
    severity_summary: pd.DataFrame,
    exception_type_summary: pd.DataFrame,
) -> None:
    """Write the stakeholder report with audit-ready tabs."""
    report_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(report_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Run Summary", index=False)
        severity_summary.to_excel(writer, sheet_name="Severity Summary", index=False)
        exception_type_summary.to_excel(writer, sheet_name="Exception Type Summary", index=False)
        exceptions.to_excel(writer, sheet_name="Exceptions", index=False)
        valid_records.to_excel(writer, sheet_name="Valid Records", index=False)

    _format_workbook(report_path)


def _format_workbook(report_path: Path) -> None:
    workbook = load_workbook(report_path)
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    money_columns = {"expected_amount", "billed_amount", "difference", "last_month_billed"}
    date_columns = {"contract_end", "invoice_date"}

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font

        headers = [cell.value for cell in worksheet[1]]

        for column_index, header in enumerate(headers, start=1):
            column_letter = get_column_letter(column_index)

            if header in money_columns:
                for cell in worksheet[column_letter][1:]:
                    cell.number_format = '"R"#,##0.00'

            if header in date_columns:
                for cell in worksheet[column_letter][1:]:
                    cell.number_format = "yyyy-mm-dd"

            max_length = len(str(header)) if header else 0
            for cell in worksheet[column_letter]:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))

            worksheet.column_dimensions[column_letter].width = min(max_length + 2, 45)

    workbook.save(report_path)
